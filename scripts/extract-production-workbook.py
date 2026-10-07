"""Read a TB production workbook without changing it; export monthly source evidence.

Usage: python scripts/extract-production-workbook.py SOURCE.xlsx OUTPUT.json
Requires openpyxl. Does not connect to a database or infer daily events.
"""
import collections
import datetime as dt
import hashlib
import json
import pathlib
import sys
import uuid
import openpyxl


def extract(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    raw = openpyxl.load_workbook(path, read_only=True, data_only=False)
    inputs = book['Inputs ']
    as_of = inputs['C19'].value
    if not isinstance(as_of, dt.datetime):
        raise ValueError('Inputs!C19 must contain the reporting date')
    data = list(book['Data Input Sheet'].values)
    model = list(book['Growth Model'].values)
    formulas = list(raw['Growth Model'].values)
    months = [(i, v) for i, v in enumerate(data[4]) if isinstance(v, dt.datetime) and v <= as_of]
    model_months = {v.date().isoformat(): i for i, v in enumerate(model[14]) if isinstance(v, dt.datetime)}
    rows = []
    for start in range(6, min(len(data), 1006), 5):
        block = data[start:start+5]
        batch = block[0][3]
        stocked_on = block[1][2]
        if not isinstance(batch, str) or not batch.strip() or not isinstance(stocked_on, dt.datetime):
            continue
        slot = int(block[0][0])
        for col, month in months:
            # Keep assigned historical observations, including zeros and source inconsistencies.
            cage = block[0][col]
            values = [block[i][col] for i in range(1, 5)]
            if not isinstance(cage, str) or cage.strip() == 'Cage ID':
                continue
            if any(v is not None and not isinstance(v, (int,float)) for v in values):
                raise ValueError(f'Invalid numeric source at row {start+1}, column {col+1}')
            mc = model_months[month.date().isoformat()]
            ms = 19 + (slot-1)*11
            source_cells = {}
            metrics = dict(zip(['stock_input_no','mortality_no','abw_end_kg','feed_kg'], values))
            for offset, key in enumerate(metrics, 1):
                source_cells[key] = f"Data Input Sheet!{openpyxl.utils.get_column_letter(col+1)}{start+offset+1}"
            calculated = {}
            for offset, key in [(2,'stock_start_no'),(3,'mortality_no'),(4,'abw_end_kg'),(5,'feed_kg'),(6,'growth_kg'),(7,'efcr'),(8,'harvest_no'),(9,'cage_correction_no'),(10,'harvest_kg')]:
                value = model[ms+offset][mc]
                formula = formulas[ms+offset][mc]
                calculated[key] = {'value': value, 'cell': f'Growth Model!{openpyxl.utils.get_column_letter(mc+1)}{ms+offset+1}', 'origin': 'calculated' if isinstance(formula,str) and formula.startswith('=') else 'entered'}
            rows.append({'source_slot':slot,'month':month.date().isoformat(),'batch_name':batch.strip(),'cage_name':cage.strip(),'stocked_on':stocked_on.date().isoformat(),**metrics,'source_cells':source_cells,'model_values':calculated})
    errors = {}
    for sheet in book:
        counts = collections.Counter(cell.value for row in sheet for cell in row if cell.data_type == 'e')
        if counts:
            errors[sheet.title] = dict(counts)
    result = {'source_name':pathlib.Path(path).name,'sha256':hashlib.sha256(pathlib.Path(path).read_bytes()).hexdigest(),'as_of':as_of.date().isoformat(),'settings':{'planned_fish':inputs['C58'].value,'planned_abw_g':inputs['C59'].value*1000,'growth_model':inputs['C14'].value,'feed_bands':[{'min_kg':inputs.cell(r,3).value,'max_kg':inputs.cell(r,4).value,'feed':inputs.cell(r,5).value} for r in range(64,70)]},'formula_errors':errors,'rows':rows}
    book.close()
    raw.close()
    return result


def import_sql(result, farm_id):
    """Generate an atomic, repeatable import; require an explicitly selected farm."""
    uuid.UUID(farm_id)
    quote = lambda value: "'" + str(value).replace("'", "''") + "'"
    fields = 'source_slot integer,month date,batch_name text,cage_name text,stocked_on date,stock_input_no numeric,mortality_no numeric,abw_end_kg numeric,feed_kg numeric,source_cells jsonb,model_values jsonb'
    columns = ','.join(part.split()[0] for part in fields.split(','))
    return f"""begin;
insert into public.production_workbook_import(farm_id,source_name,sha256,as_of,settings,formula_errors)
values ({quote(farm_id)},{quote(result['source_name'])},{quote(result['sha256'])},{quote(result['as_of'])},{quote(json.dumps(result['settings']))}::jsonb,{quote(json.dumps(result['formula_errors']))}::jsonb)
on conflict(farm_id,sha256) do nothing;
insert into public.production_workbook_month(import_id,farm_id,{columns})
select i.id,i.farm_id,{','.join('r.'+c for c in columns.split(','))}
from public.production_workbook_import i cross join jsonb_to_recordset({quote(json.dumps(result['rows']))}::jsonb) as r({fields})
where i.farm_id={quote(farm_id)} and i.sha256={quote(result['sha256'])}
on conflict(import_id,source_slot,month) do nothing;
commit;
"""


if __name__ == '__main__':
    result = extract(sys.argv[1])
    pathlib.Path(sys.argv[2]).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    if len(sys.argv) == 5:
        pathlib.Path(sys.argv[4]).write_text(import_sql(result, sys.argv[3]), encoding='utf-8')
    print(json.dumps({'rows':len(result['rows']),'months':sorted(set(r['month'] for r in result['rows'])),'batches':sorted(set(r['batch_name'] for r in result['rows'])),'sha256':result['sha256']}))
